# -*- coding: utf-8 -*-
"""
桌面课表工具 · 配置程序（tkinter）
功能：
  1. 编辑单周课表 / 双周课表（节次 × 星期 下拉选择科目）；
  2. 临时调整任意一天课程（优先于正常课表，主程序当日自动生效），可填备注；
  3. 课程管理：新增学科（对话框 + 调色盘选色）、修改颜色、删除学科（同时从所有课表中移除）；
  4. 设置学期起始日（用于推算单双周），支持强制单周/双周；
  5. 每节课时间段（可选）、主程序显示设置（缩放/位置）；
  6. 导入 / 导出课表（JSON）、从 Excel 导入课表（自动补未知科目）；
  7. 管理"开机自动启动主程序"。

所有修改实时保存到 schedule_data.json（与主程序同目录）。

用法：
  python config_app.py            启动配置程序
  python config_app.py --selftest 构造界面后自动保存并退出（调试用）

适用：Python 3.13，Windows。
"""

import argparse
import ctypes
import datetime
import json
import os
import subprocess
import sys
import tkinter as tk
from tkinter import colorchooser, filedialog, messagebox, ttk

import schedule_common as sc


def read_schedule_from_file(path):
    """从 Excel（.xlsx/.xls）或 CSV 读取课表，返回 {星期: [课程...]}。

    约定格式：第一行为表头（须包含"周一"~"周日"之一或多个），
    其余每行一节课，对应列即为该星期的课程；空单元格表示无课。
    """
    import csv
    ext = os.path.splitext(path)[1].lower()
    rows = []
    if ext == ".csv":
        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            rows = [list(r) for r in csv.reader(f)]
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        rows = [[("" if c is None else str(c).strip()) for c in r]
                for r in ws.iter_rows(values_only=True)]
    elif ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        sh = book.sheet_by_index(0)
        rows = [[("" if sh.cell_value(r, c) == "" else str(sh.cell_value(r, c)).strip())
                 for c in range(sh.ncols)] for r in range(sh.nrows)]
    else:
        raise ValueError("不支持的文件格式：%s（请使用 .xlsx / .xls / .csv）" % ext)

    if not rows:
        raise ValueError("文件内容为空")
    day_idx = {}
    for i, cell in enumerate(rows[0]):
        name = (cell or "").strip()
        if name in sc.WEEKDAYS:
            day_idx[i] = name
    if not day_idx:
        raise ValueError("未在首行找到星期表头（周一~周日）")
    result = {d: [] for d in sc.WEEKDAYS}
    for row in rows[1:]:
        for i, d in day_idx.items():
            result[d].append(row[i].strip() if i < len(row) else "")
    return result


class ConfigApp:
    def __init__(self, root):
        self.root = root
        self.cfg = sc.load_config()
        self.grids = {}
        self.temp_panel = None
        self.settings = None
        self.status_var = tk.StringVar()

        root.title("桌面课表 · 配置程序")
        root.geometry("1020x680")
        root.minsize(900, 560)
        root.protocol("WM_DELETE_WINDOW", self._on_close)

        self._build()
        self.refresh_all()
        self.status("就绪：配置已加载（所有修改自动保存到 schedule_data.json）")

    # ---------------- 界面 ----------------
    def _build(self):
        self.nb = ttk.Notebook(self.root)
        self.nb.pack(fill="both", expand=True, padx=6, pady=(6, 0))

        t1 = ttk.Frame(self.nb)
        t2 = ttk.Frame(self.nb)
        t3 = ttk.Frame(self.nb)
        t4 = ttk.Frame(self.nb)
        t5 = ttk.Frame(self.nb)
        self.nb.add(t1, text="  单周课表  ")
        self.nb.add(t2, text="  双周课表  ")
        self.nb.add(t3, text="  临时调课  ")
        self.nb.add(t4, text="  课程管理  ")
        self.nb.add(t5, text="  设置  ")

        self.grids["single"] = ScheduleGrid(t1, self, "single")
        self.grids["double"] = ScheduleGrid(t2, self, "double")
        self.grids["single"].pack(fill="both", expand=True)
        self.grids["double"].pack(fill="both", expand=True)

        self.temp_panel = TempAdjustPanel(t3, self)
        self.temp_panel.pack(fill="both", expand=True)

        self.subjects_panel = SubjectManagerPanel(t4, self)
        self.subjects_panel.pack(fill="both", expand=True)

        self.settings = SettingsPanel(t5, self)
        self.settings.pack(fill="both", expand=True)

        bar = ttk.Frame(self.root)
        bar.pack(fill="x", padx=6, pady=4)
        for text, cmd in (("保存设置", self.save),
                          ("导入课表…", self.import_config),
                          ("从Excel导入…", self.import_from_excel),
                          ("导出课表…", self.export_config),
                          ("打开配置目录", self.open_dir)):
            ttk.Button(bar, text=text, command=cmd).pack(side="left", padx=2)
        ttk.Label(bar, textvariable=self.status_var).pack(side="right", padx=6)

    # ---------------- 数据操作 ----------------
    def save(self, msg="配置已保存"):
        sc.save_config(self.cfg)
        self.status(msg)

    def status(self, text):
        self.status_var.set(text)

    def refresh_all(self):
        for g in self.grids.values():
            g.render_rows()
        self.temp_panel.render_rows()
        self.settings.refresh_preview()
        self.subjects_panel.refresh_list()

    def _on_close(self):
        sc.save_config(self.cfg)
        self.root.destroy()

    def import_config(self):
        path = filedialog.askopenfilename(
            title="导入课表配置",
            filetypes=[("JSON 配置", "*.json"), ("所有文件", "*.*")])
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as e:
            messagebox.showerror("导入失败", "无法读取文件：\n%s" % e)
            return
        if not isinstance(data, dict):
            messagebox.showerror("导入失败", "文件内容不是 JSON 对象（课表配置）。")
            return
        sc.deep_merge(self.cfg, data)
        if not self.cfg["subjects"]:
            self.cfg["subjects"] = dict(sc.DEFAULT_SUBJECTS)
        sc.save_config(self.cfg)
        self.refresh_all()
        self.status("已导入：" + os.path.basename(path))

    def import_from_excel(self):
        """从 Excel/CSV 导入课表到单周或双周。"""
        path = filedialog.askopenfilename(
            title="从 Excel 导入课表",
            filetypes=[("Excel/CSV", "*.xlsx *.xls *.csv"), ("所有文件", "*.*")])
        if not path:
            return
        try:
            data = read_schedule_from_file(path)
        except Exception as e:
            messagebox.showerror("导入失败", "%s" % e)
            return
        which = messagebox.askyesno(
            "导入目标",
            "将课表导入到【单周课表】？\n（选择“否”则导入到【双周课表】）")
        key = "single" if which else "double"

        # 未知科目自动加入科目表
        added = []
        for d in sc.WEEKDAYS:
            for name in data[d]:
                if name and name not in self.cfg["subjects"]:
                    self.cfg["subjects"][name] = {
                        "color": "#95A5A6", "short": name[0]}
                    added.append(name)

        self.cfg["schedule"][key] = data
        sc.save_config(self.cfg)
        self.refresh_all()
        self.subjects_panel.refresh_list()
        msg = "已从 %s 导入到【%s课表】" % (
            os.path.basename(path), "单周" if which else "双周")
        if added:
            msg += "；自动新增科目：" + "、".join(dict.fromkeys(added))
        self.status(msg)

    def export_config(self):
        path = filedialog.asksaveasfilename(
            title="导出课表配置", defaultextension=".json",
            initialfile="schedule_data.json",
            filetypes=[("JSON 配置", "*.json")])
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(self.cfg, f, ensure_ascii=False, indent=2)
        except Exception as e:
            messagebox.showerror("导出失败", str(e))
            return
        self.status("已导出：" + path)

    def open_dir(self):
        os.startfile(sc.APP_DIR)


class ScheduleGrid(ttk.Frame):
    """单周 / 双周课表编辑器：行=节次，列=星期，下拉选择科目。"""

    def __init__(self, master, app, week_key):
        super().__init__(master, padding=6)
        self.app = app
        self.week_key = week_key
        self.vars = []

        bar = ttk.Frame(self)
        bar.pack(fill="x", pady=(0, 4))
        ttk.Label(bar, text="每一行是一节课，每一列是一个星期；下拉选择科目（“（空）”表示无课）。"
                            ).pack(side="left")
        ttk.Button(bar, text="＋ 增加一节", command=self.add_row).pack(side="right", padx=2)
        ttk.Button(bar, text="－ 删除末节", command=self.remove_last).pack(side="right", padx=2)

        self.host = ttk.Frame(self)
        self.host.pack(fill="both", expand=True)
        self.render_rows()

    def schedule(self):
        return self.app.cfg["schedule"][self.week_key]

    def render_rows(self):
        for w in self.host.winfo_children():
            w.destroy()
        subjects = list(self.app.cfg["subjects"].keys())
        choices = ["（空）"] + subjects

        hdr = ttk.Frame(self.host)
        hdr.pack(fill="x")
        ttk.Label(hdr, text="节次", width=10).pack(side="left")
        for d in sc.WEEKDAYS:
            ttk.Label(hdr, text=d, width=11).pack(side="left")

        sched = self.schedule()
        periods = max((len(sched.get(d, [])) for d in sc.WEEKDAYS), default=0)
        # 默认显示 default_periods 行（未填满的以"（空）"补足）
        periods = max(periods, int(self.app.cfg.get("default_periods", 10)))
        self.vars = []
        for i in range(periods):
            row = ttk.Frame(self.host)
            row.pack(fill="x", pady=1)
            ttk.Label(row, text="第%d节" % (i + 1), width=10).pack(side="left")
            row_vars = []
            for d in sc.WEEKDAYS:
                lst = sched.get(d, [])
                raw = lst[i] if i < len(lst) else ""
                v = tk.StringVar(value=raw or "（空）")
                cb = ttk.Combobox(row, textvariable=v, values=choices,
                                  width=10, state="readonly")
                cb.pack(side="left", padx=1)
                v.trace_add("write", lambda *_, vv=v: self._commit())
                row_vars.append(v)
            self.vars.append(row_vars)

    def _commit(self):
        sched = self.schedule()
        for i, row in enumerate(self.vars):
            for j, v in enumerate(row):
                lst = sched.setdefault(sc.WEEKDAYS[j], [])
                while len(lst) <= i:
                    lst.append("")
                val = v.get()
                lst[i] = "" if val == "（空）" else val
        self.app.save()

    def add_row(self):
        self._commit()
        n = len(self.vars)
        for d in sc.WEEKDAYS:
            lst = self.schedule().setdefault(d, [])
            if len(lst) <= n:
                lst.append("")
        self.render_rows()
        self.app.save()

    def remove_last(self):
        if len(self.vars) <= 1:
            return
        self._commit()
        for d in sc.WEEKDAYS:
            lst = self.schedule().get(d, [])
            if lst:
                lst.pop()
        self.render_rows()
        self.app.save()


class TempAdjustPanel(ttk.Frame):
    """临时调课：指定日期 + 课程列表 + 备注，优先于正常课表。"""

    def __init__(self, master, app):
        super().__init__(master, padding=6)
        self.app = app
        self.date_var = tk.StringVar(value=datetime.date.today().strftime("%Y-%m-%d"))
        self.note_var = tk.StringVar()
        self.period_vars = []
        self._force_normal = False

        top = ttk.Frame(self)
        top.pack(fill="x")
        ttk.Label(top, text="调整日期（YYYY-MM-DD）：").pack(side="left")
        ttk.Entry(top, textvariable=self.date_var, width=12).pack(side="left", padx=4)
        ttk.Button(top, text="载入当天正常课表", command=self.load_normal).pack(side="left", padx=2)
        ttk.Button(top, text="保存调整", command=self.save_adjust).pack(side="left", padx=2)
        ttk.Button(top, text="清除该日调整", command=self.clear_adjust).pack(side="left", padx=2)
        ttk.Label(top, text="调整优先于正常课表，主程序当日自动生效。",
                  foreground="#666666").pack(side="left", padx=8)

        nrow = ttk.Frame(self)
        nrow.pack(fill="x", pady=4)
        ttk.Label(nrow, text="备注（如：运动会调课）：").pack(side="left")
        ttk.Entry(nrow, textvariable=self.note_var, width=44).pack(side="left", padx=4)

        self.host = ttk.Frame(self)
        self.host.pack(fill="both", expand=True, pady=4)
        self.hint_var = tk.StringVar()
        ttk.Label(self, textvariable=self.hint_var, foreground="#555555").pack(anchor="w")

        self.render_rows()

    # ---------------- 逻辑 ----------------
    def _date(self):
        try:
            return datetime.datetime.strptime(self.date_var.get().strip(), "%Y-%m-%d").date()
        except ValueError:
            messagebox.showwarning("日期格式错误", "日期格式应为 YYYY-MM-DD，例如 2026-09-27")
            return None

    def normal_courses(self, d):
        parity, week = sc.week_parity(d, self.app.cfg)
        key = "single" if parity == "single" else "double"
        lst = self.app.cfg["schedule"].get(key, {}).get(sc.WEEKDAYS[d.weekday()], [])
        return list(lst), parity, week

    def render_rows(self):
        for w in self.host.winfo_children():
            w.destroy()
        d = self._date()
        if d is None:
            return
        existing = self.app.cfg["temp_adjustments"].get(d.strftime("%Y-%m-%d"))
        courses, parity, week = self.normal_courses(d)

        if existing and not self._force_normal:
            courses = list(existing.get("periods") or [])
            self.note_var.set(existing.get("note") or "")
        else:
            self.note_var.set("")

        subjects = list(self.app.cfg["subjects"].keys())
        choices = ["（空）"] + subjects
        if not courses:
            courses = [""]

        hdr = ttk.Frame(self.host)
        hdr.pack(fill="x")
        ttk.Label(hdr, text="节次", width=10).pack(side="left")
        ttk.Label(hdr, text="课程", width=16).pack(side="left")
        ttk.Label(hdr, text="（节数可与正常课表不同，以本列表为准）",
                  foreground="#888888").pack(side="left", padx=6)

        self.period_vars = []
        for i, name in enumerate(courses, start=1):
            row = ttk.Frame(self.host)
            row.pack(fill="x", pady=1)
            ttk.Label(row, text="第%d节" % i, width=10).pack(side="left")
            v = tk.StringVar(value=name or "（空）")
            ttk.Combobox(row, textvariable=v, values=choices, width=16,
                         state="readonly").pack(side="left")
            self.period_vars.append(v)

        if existing and not self._force_normal:
            self.hint_var.set("%s 已存在手动调整（将优先生效）" % d)
        else:
            self.hint_var.set("%s（%s）：正常课表为第%s周 · %s，可修改后点击“保存调整”。"
                              % (d, sc.WEEKDAYS[d.weekday()], week,
                                 "单周" if parity == "single" else "双周"))

    def load_normal(self):
        d = self._date()
        if not d:
            return
        self._force_normal = True
        self.note_var.set("")
        self.render_rows()
        self._force_normal = False
        self.app.status("已载入 %s 正常课表，可修改后点击“保存调整”" % d)

    def save_adjust(self):
        d = self._date()
        if not d:
            return
        courses = ["" if v.get() == "（空）" else v.get() for v in self.period_vars]
        while courses and courses[-1] == "":
            courses.pop()
        self.app.cfg["temp_adjustments"][d.strftime("%Y-%m-%d")] = {
            "periods": courses,
            "note": self.note_var.get().strip(),
        }
        self.app.save("已保存 %s 的临时调整" % d)
        self.render_rows()

    def clear_adjust(self):
        d = self._date()
        if not d:
            return
        key = d.strftime("%Y-%m-%d")
        if key in self.app.cfg["temp_adjustments"]:
            del self.app.cfg["temp_adjustments"][key]
        self.note_var.set("")
        self.app.save("已清除 %s 的临时调整" % d)
        self.render_rows()


class SubjectManagerPanel(ttk.Frame):
    """课程管理：查看 / 新增 / 删除学科、修改颜色。"""

    def __init__(self, master, app):
        super().__init__(master, padding=8)
        self.app = app
        self.names = []
        self._build()

    def _build(self):
        tip = ttk.Label(self, text="管理学科列表：可新增、修改颜色或删除学科；"
                                   "删除学科时会同时从单双周课表与临时调课中移除该课程。",
                        foreground="#555555")
        tip.pack(anchor="w", pady=(0, 6))

        host = ttk.Frame(self)
        host.pack(fill="both", expand=True)
        self.subj_list = tk.Listbox(host, width=34, height=14)
        self.subj_list.pack(side="left", fill="both", expand=True, padx=(0, 8))
        sb = ttk.Scrollbar(host, orient="vertical", command=self.subj_list.yview)
        sb.pack(side="left", fill="y")
        self.subj_list.configure(yscrollcommand=sb.set)
        self.subj_list.bind("<Double-Button-1>", lambda e: self.pick_color())
        self.subj_list.bind("<<ListboxSelect>>", self._on_select)

        self.swatch = tk.Label(host, text="  ", bg="#FFFFFF", width=4, height=3,
                               relief="solid", borderwidth=1)
        self.swatch.pack(side="left", padx=8)
        info = ttk.Frame(host)
        info.pack(side="left", fill="y")
        ttk.Label(info, text="当前选中：").pack(anchor="w")
        self.sel_name_var = tk.StringVar(value="（未选中）")
        ttk.Label(info, textvariable=self.sel_name_var,
                  foreground="#333333").pack(anchor="w", pady=(2, 4))
        ttk.Label(info, text="（双击列表项也可直接修改颜色）",
                  foreground="#888888").pack(anchor="w")

        btns = ttk.Frame(self)
        btns.pack(fill="x", pady=(8, 0))
        ttk.Button(btns, text="修改颜色…", command=self.pick_color).pack(side="left", padx=2)
        ttk.Button(btns, text="新增学科…", command=self.add_subject).pack(side="left", padx=2)
        ttk.Button(btns, text="删除学科…", command=self.delete_subject).pack(side="left", padx=2)
        ttk.Button(btns, text="恢复默认颜色", command=self.reset_defaults).pack(side="left", padx=2)

        self.refresh_list()

    def refresh_list(self):
        """按当前配置刷新科目列表与选中预览。"""
        self.subj_list.delete(0, "end")
        self.names = list(self.app.cfg["subjects"].keys())
        for name in self.names:
            info = self.app.cfg["subjects"][name]
            self.subj_list.insert("end", "%s  %s" % (name, info.get("color", "")))
        self._on_select()

    def _on_select(self, event=None):
        sel = self.subj_list.curselection()
        if sel:
            name = self.names[sel[0]]
            self.sel_name_var.set(name)
            self.swatch.configure(bg=self.app.cfg["subjects"][name].get("color", "#888888"))
        else:
            self.sel_name_var.set("（未选中）")
            self.swatch.configure(bg="#FFFFFF")

    def pick_color(self):
        sel = self.subj_list.curselection()
        if not sel:
            messagebox.showinfo("提示", "请先在列表中选中一个科目")
            return
        name = self.names[sel[0]]
        cur = self.app.cfg["subjects"][name].get("color", "#888888")
        rgb, hexv = colorchooser.askcolor(color=cur, title="%s 颜色" % name)
        if hexv:
            self.app.cfg["subjects"][name]["color"] = hexv
            self.app.save("已更新 %s 颜色" % name)
            self.refresh_list()
            self.subj_list.selection_set(sel[0])

    def add_subject(self):
        AddSubjectDialog(self, self.app)

    def delete_subject(self):
        sel = self.subj_list.curselection()
        if not sel:
            messagebox.showinfo("提示", "请先在列表中选中一个要删除的科目")
            return
        name = self.names[sel[0]]
        if not messagebox.askyesno(
                "删除学科",
                "确定删除学科【%s】吗？\n将同时从单双周课表与临时调课中移除该课程。" % name):
            return
        del self.app.cfg["subjects"][name]
        # 从单双周课表所有星期中移除该课程
        for wk in ("single", "double"):
            for d in sc.WEEKDAYS:
                lst = self.app.cfg["schedule"].get(wk, {}).get(d, [])
                self.app.cfg["schedule"][wk][d] = [c for c in lst if c != name]
        # 从临时调课中移除；若某日调整因此为空则删除该日调整
        for day, adj in list(self.app.cfg["temp_adjustments"].items()):
            periods = [c for c in adj.get("periods", []) if c != name]
            if not periods:
                del self.app.cfg["temp_adjustments"][day]
            else:
                adj["periods"] = periods
        self.app.save("已删除学科：" + name)
        self.app.refresh_all()
        self.refresh_list()

    def reset_defaults(self):
        if messagebox.askyesno(
                "恢复默认颜色",
                "将内置 16 个科目的颜色恢复为默认值？\n（自定义新增的学科不会被删除）"):
            for name, info in dict(sc.DEFAULT_SUBJECTS).items():
                if name in self.app.cfg["subjects"]:
                    self.app.cfg["subjects"][name].update(info)
            self.app.save("科目颜色已恢复默认")
            self.refresh_list()


class SettingsPanel(ttk.Frame):
    """学期起始日 / 单双周模式 / 时间段 / 主程序显示 / 开机自启。"""

    def __init__(self, master, app):
        super().__init__(master, padding=8)
        self.app = app
        self._build()

    def _build(self):
        left = ttk.LabelFrame(self, text=" 学期与单双周 ", padding=8)
        left.pack(side="left", fill="y", padx=(0, 8))
        ttk.Label(left, text="学期起始日（建议填开学第一周的周一）：").pack(anchor="w")
        row = ttk.Frame(left)
        row.pack(fill="x", pady=4)
        self.start_var = tk.StringVar(value=self.app.cfg["semester_start_date"] or "")
        ttk.Entry(row, textvariable=self.start_var, width=12).pack(side="left")
        ttk.Button(row, text="设为本周周一", command=self.set_start_this_week).pack(side="left", padx=4)

        self.preview_var = tk.StringVar()
        ttk.Label(left, textvariable=self.preview_var,
                  foreground="#c0392b").pack(anchor="w", pady=(4, 8))
        ttk.Separator(left).pack(fill="x", pady=4)
        ttk.Label(left, text="单双周模式：").pack(anchor="w")
        self.mode_var = tk.StringVar(value=self._mode_key())
        for text, val in (("自动（按学期起始日推算）", "auto"),
                          ("强制单周", "single"),
                          ("强制双周", "double")):
            ttk.Radiobutton(left, text=text, value=val, variable=self.mode_var,
                            command=self.apply_mode).pack(anchor="w")
        ttk.Button(left, text="刷新推算", command=self.refresh_preview).pack(anchor="w", pady=(6, 0))

        right = ttk.LabelFrame(self, text=" 每节课时间（可选，主程序显示用） ", padding=8)
        right.pack(side="left", fill="both", expand=True)
        self.time_text = tk.Text(right, width=30, height=8)
        self.time_text.pack(fill="x")
        ttk.Label(right, text="每行一节课，格式：08:00-08:45",
                  foreground="#666666").pack(anchor="w")
        ttk.Button(right, text="应用时间段", command=self.apply_times).pack(anchor="w", pady=4)
        ttk.Separator(right).pack(fill="x", pady=6)

        disp = ttk.LabelFrame(right, text=" 主程序显示 ", padding=6)
        disp.pack(fill="x", pady=(0, 6))
        win = self.app.cfg.get("window", {})
        self.scale_var = tk.BooleanVar(value=bool(win.get("scale_to_screen")))
        ttk.Checkbutton(disp, text="按屏幕分辨率自动缩放\n（右侧：高度占屏高 4/5；顶部：宽度占屏宽 3/4）",
                        variable=self.scale_var, command=self.apply_window).pack(anchor="w")
        self.pos_var = tk.StringVar(value=win.get("mode", "right"))
        ttk.Radiobutton(disp, text="位于屏幕右侧（竖排）", value="right",
                        variable=self.pos_var, command=self.apply_window).pack(anchor="w")
        ttk.Radiobutton(disp, text="位于屏幕顶部（横排）", value="top",
                        variable=self.pos_var, command=self.apply_window).pack(anchor="w")
        ttk.Separator(right).pack(fill="x", pady=4)

        self.auto_var = tk.BooleanVar(value=sc.is_autostart_installed())
        ttk.Checkbutton(right, text="开机自动启动主程序（当前用户）",
                        variable=self.auto_var,
                        command=self.toggle_autostart).pack(anchor="w")
        ttk.Button(right, text="立即启动主程序", command=self.launch_main).pack(anchor="w", pady=(6, 0))
        ttk.Button(right, text="恢复默认课表（清空编辑）",
                   command=self.reset_all).pack(anchor="w", pady=(4, 0))

        self.time_text.delete("1.0", "end")
        self.time_text.insert("1.0", "\n".join(self.app.cfg.get("period_times", [])))

    # ---------------- 逻辑 ----------------
    def _mode_key(self):
        o = self.app.cfg.get("week_override")
        return o if o in ("single", "double") else "auto"

    def apply_mode(self):
        m = self.mode_var.get()
        self.app.cfg["week_override"] = None if m == "auto" else m
        self.app.save("单双周模式已更新")
        self.refresh_preview()

    def set_start_this_week(self):
        today = datetime.date.today()
        monday = today - datetime.timedelta(days=today.weekday())
        self.start_var.set(monday.strftime("%Y-%m-%d"))
        self.refresh_preview()

    def refresh_preview(self):
        s = self.start_var.get().strip()
        if s:
            try:
                datetime.datetime.strptime(s, "%Y-%m-%d")
            except ValueError:
                self.preview_var.set("起始日格式错误：应为 YYYY-MM-DD")
                return
            if self.app.cfg.get("semester_start_date") != s:
                self.app.cfg["semester_start_date"] = s
                self.app.save("学期起始日已更新")
        today = datetime.date.today()
        parity, week = sc.week_parity(today, self.app.cfg)
        mode = {"auto": "自动", "single": "强制单周", "double": "强制双周"}[self.mode_var.get()]
        self.preview_var.set("今日：%s %s → 第%s周·%s（模式：%s）"
                             % (today, sc.WEEKDAYS[today.weekday()], week,
                                "单周" if parity == "single" else "双周", mode))

    def apply_times(self):
        lines = [ln.strip() for ln in self.time_text.get("1.0", "end").splitlines() if ln.strip()]
        self.app.cfg["period_times"] = lines
        self.app.save("每节课时间已更新（%d 节）" % len(lines))

    def toggle_autostart(self):
        if self.auto_var.get():
            cmd = sc.install_autostart()
            self.app.status("已开启开机自启：" + cmd)
        else:
            sc.remove_autostart()
            self.app.status("已取消开机自启")

    def apply_window(self):
        win = self.app.cfg.setdefault("window", {})
        win["scale_to_screen"] = bool(self.scale_var.get())
        win["mode"] = self.pos_var.get()
        self.app.save("主程序显示设置已更新（主程序自动生效）")

    def launch_main(self):
        if getattr(sys, "frozen", False):
            subprocess.Popen([os.path.join(sc.APP_DIR, "桌面课表.exe")])
        else:
            subprocess.Popen([sys.executable, os.path.join(sc.APP_DIR, "main_app.py")])
        self.app.status("主程序已启动")

    def reset_all(self):
        if messagebox.askyesno("恢复默认", "将清空单双周课表与临时调整并恢复默认设置，继续？"):
            self.app.cfg = sc.load_defaults()
            self.app.save("已恢复默认配置")
            self.app.refresh_all()


class AddSubjectDialog(tk.Toplevel):
    """对话框式新增学科：输入名称 + 调色盘选择颜色。"""

    def __init__(self, parent, app):
        super().__init__(parent)
        self.app = app
        self.title("新增学科")
        self.resizable(False, False)
        self.transient(parent.winfo_toplevel())
        self.grab_set()

        body = ttk.Frame(self, padding=14)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="学科名称：").grid(row=0, column=0, sticky="w", pady=4)
        self.name_var = tk.StringVar()
        ttk.Entry(body, textvariable=self.name_var, width=18).grid(
            row=0, column=1, sticky="w", padx=6)
        ttk.Label(body, text="图标颜色：").grid(row=1, column=0, sticky="w", pady=4)
        self.color_var = tk.StringVar(value="#95A5A6")
        self.color_btn = tk.Button(body, text="", width=6, bg=self.color_var.get(),
                                   command=self.pick)
        self.color_btn.grid(row=1, column=1, sticky="w", padx=6)
        ttk.Label(body, text="（点击色块打开调色盘）",
                  foreground="#888888").grid(row=1, column=2, sticky="w")

        btns = ttk.Frame(body)
        btns.grid(row=2, column=0, columnspan=3, pady=(12, 0))
        ttk.Button(btns, text="确定", command=self.ok).pack(side="left", padx=8)
        ttk.Button(btns, text="取消", command=self.destroy).pack(side="left")

        self.bind("<Return>", lambda e: self.ok())
        self.bind("<Escape>", lambda e: self.destroy())

    def pick(self):
        rgb, hexv = colorchooser.askcolor(color=self.color_var.get(), parent=self,
                                          title="选择颜色")
        if hexv:
            self.color_var.set(hexv)
            self.color_btn.configure(bg=hexv)

    def ok(self):
        name = self.name_var.get().strip()
        if not name:
            messagebox.showwarning("提示", "请填写学科名称", parent=self)
            return
        if name in self.app.cfg["subjects"]:
            messagebox.showinfo("提示", "该学科已存在", parent=self)
            return
        self.app.cfg["subjects"][name] = {
            "color": self.color_var.get(), "short": name[0]}
        self.app.save("已新增学科：" + name)
        self.app.refresh_all()
        self.app.subjects_panel.refresh_list()
        self.destroy()


def selftest():
    """构造完整界面但不进入主循环，验证 UI 构建与保存链路。"""
    root = tk.Tk()
    root.withdraw()
    app = ConfigApp(root)
    app.save("selftest saved")
    root.update_idletasks()
    root.destroy()
    print("CONFIG_SELFTEST_OK")


def main():
    ap = argparse.ArgumentParser(description="桌面课表配置程序")
    ap.add_argument("--selftest", action="store_true", help="构造界面后自动保存并退出")
    args = ap.parse_args()
    if args.selftest:
        selftest()
        return
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    root = tk.Tk()
    ConfigApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
